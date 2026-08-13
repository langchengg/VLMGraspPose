#!/usr/bin/env python3
"""Build the bilingual 16:9 case-analysis deck with python-pptx."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


BLUE = RGBColor(0, 114, 178)
ORANGE = RGBColor(230, 159, 0)
GREEN = RGBColor(0, 158, 115)
MAGENTA = RGBColor(204, 121, 167)
DARK = RGBColor(32, 38, 46)
LIGHT = RGBColor(246, 248, 250)


def _blank(prs: Presentation, title: str, subtitle: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = LIGHT
    box = slide.shapes.add_textbox(
        Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.7)
    )
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = DARK
    if subtitle:
        sub = slide.shapes.add_textbox(
            Inches(0.58), Inches(0.86), Inches(12), Inches(0.35)
        )
        p = sub.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(85, 92, 101)
    return slide


def _image(
    slide, path: Path, left: float, top: float, width: float, height: float
) -> None:
    with Image.open(path) as image:
        ratio = image.width / image.height
    target = width / height
    if ratio >= target:
        shown_w = width
        shown_h = width / ratio
        x = left
        y = top + (height - shown_h) / 2
    else:
        shown_h = height
        shown_w = height * ratio
        x = left + (width - shown_w) / 2
        y = top
    slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), width=Inches(shown_w), height=Inches(shown_h)
    )


def _bullets(
    slide, items: list[str], left=0.7, top=1.35, width=12, height=5.5, size=19
) -> None:
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(12)


def build(run: Path) -> Path:
    run = run.resolve()
    manifest = json.loads((run / "manifest.json").read_text())
    performance = pd.read_csv(run / "tables/formal_reconciliation.csv")
    gallery = pd.read_csv(run / "05_case_selection/core_case_selection.csv")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    notes: list[tuple[str, str]] = []

    slide = _blank(
        prs,
        "Frozen reranking: success, failure, and mechanism",
        "锁定后只读分析 · CROG / HiFi-CS→G1 / HiFi-CS→C1",
    )
    _bullets(
        slide,
        [
            "No retraining, reselection, gate tuning, or formal-Test rerun",
            "23,025 sample-route decisions · 9 frozen LightGBM models · exact score replay",
            f"Source final lock: {manifest['source_final_lock_sha256'][:16]}…",
        ],
        top=1.8,
        size=23,
    )
    notes.append(
        (
            "Title",
            "Open with the evidence boundary: this is a derived post-formal analysis, not a new formal result.",
        )
    )

    slide = _blank(prs, "1 · Evidence contract / 证据边界")
    _bullets(
        slide,
        [
            "Formal outcome: locked candidate_success and one-shot Test decisions",
            "Validation evidence: predeclared cumulative and leave-one-family-out ablations",
            "Model evidence: native LightGBM additive score contributions (post-hoc, non-causal)",
            "Human evidence: deterministic case-board interpretation",
            "Offline 2D Jacquard consistency ≠ physical grasp success",
        ],
    )
    notes.append(("Evidence", "Keep the four evidence layers separate in discussion."))

    for number, stem, title in (
        (2, "01_native_ungated_gated_j1", "Exact formal reconciliation"),
        (3, "02_oracle_headroom", "How much order-only headroom existed?"),
        (4, "03_recovery_harm_waterfall", "Recovered versus harmful switches"),
        (5, "04_positive_rate_by_native_rank", "Where correct candidates occur"),
        (6, "05_rank_transition_heatmap", "How the learned order changes"),
        (7, "07_gain_importance", "Frozen-tree global importance"),
        (
            8,
            "09_recovered_harmful_family_contributions",
            "Recovered and harmful mechanisms",
        ),
        (9, "11_gate_margin_utility", "Conservative-gate boundary"),
        (10, "13_irreparable_bottleneck", "Why G1/C1 remain below CROG"),
        (
            11,
            "15_validation_ablation_triangulation",
            "Validation ablation triangulation",
        ),
    ):
        slide = _blank(prs, f"{number} · {title}")
        _image(slide, run / "08_figures" / f"{stem}.png", 0.65, 1.15, 12.05, 5.9)
        notes.append(
            (
                title,
                "Read the plot as association and frozen formal outcome; avoid causal language.",
            )
        )

    for number, route, category, title in (
        (12, "crog", "recovered", "CROG recovered case"),
        (13, "g1", "recovered", "G1 recovered case"),
        (14, "c1", "harmful", "C1 harmful switch"),
    ):
        sub = gallery[(gallery["route"] == route) & (gallery["category"] == category)]
        if sub.empty:
            sub = gallery[gallery["route"] == route].head(1)
        slide = _blank(
            prs,
            f"{number} · {title}",
            "cyan=native · orange=ungated challenger · magenta=final · blue dashed=GT",
        )
        _image(slide, Path(sub.iloc[0]["board_path"]), 0.45, 1.15, 12.45, 5.95)
        notes.append(
            (
                title,
                "Explain the decision chain from native Top-1 to challenger to the gated final candidate.",
            )
        )

    slide = _blank(prs, "15 · Conclusions / 结论")
    p = performance.set_index("route")
    _bullets(
        slide,
        [
            f"CROG: {p.loc['crog', 'native_j_at_1']:.3%} → {p.loc['crog', 'gated_j_at_1']:.3%}; high starting quality, limited remaining Top-5 headroom",
            f"G1: {p.loc['g1', 'native_j_at_1']:.3%} → {p.loc['g1', 'gated_j_at_1']:.3%}; 3,206 reranking-irreparable samples",
            f"C1: {p.loc['c1', 'native_j_at_1']:.3%} → {p.loc['c1', 'gated_j_at_1']:.3%}; 3,163 reranking-irreparable samples",
            "Matched-common evidence changes order effectively; the gate reduces harmful switches",
            "Next bottleneck is candidate generation/grounded evidence—not only ranking",
        ],
        top=1.35,
        size=18,
    )
    notes.append(
        (
            "Conclusions",
            "Use the exact scientific name: CROG candidate pool with external matched RGB-D/HiFi evidence reranking.",
        )
    )

    output = run / "10_slides/RERANKING_SUCCESS_FAILURE_ANALYSIS.pptx"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    (run / "10_slides/SPEAKER_NOTES.md").write_text(
        "# Speaker notes\n\n"
        + "\n\n".join(
            f"## Slide {i + 1}: {title}\n\n{text}"
            for i, (title, text) in enumerate(notes)
        )
        + "\n"
    )
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", required=True, type=Path)
    print(build(parser.parse_args().run_dir))


if __name__ == "__main__":
    main()
