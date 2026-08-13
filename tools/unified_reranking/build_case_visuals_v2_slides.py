#!/usr/bin/env python3
"""Build full and short 16:9 decks from verified v2 case-board assets."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROUTES = ("crog", "g1", "c1")
DISPLAY = {"crog": "CROG", "g1": "HiFi-CS → G1", "c1": "HiFi-CS → C1"}
AUDIT_ORDER = (
    "recovered",
    "harmful",
    "gate_prevented_harmful",
    "gate_missed_recoverable",
    "wrong_to_wrong_solvable",
    "candidate_generation_irreparable",
)
W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(21, 32, 43)
MUTED = RGBColor(85, 97, 111)
BLUE = RGBColor(0, 87, 184)
MAGENTA = RGBColor(204, 121, 167)


def _new() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def _textbox(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float,
    bold: bool = False,
    color: RGBColor = INK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _title_slide(prs: Presentation, *, short: bool) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(245, 247, 250)
    _textbox(slide, "Three-route reranking: success, failure, and earliest observable module", 0.65, 0.72, 12.0, 1.2, size=32, bold=True)
    _textbox(slide, "Presentation-selected cases from the complete post-formal Test population", 0.67, 1.83, 11.8, 0.55, size=20, color=MUTED)
    metrics = [
        ("CROG", "89.225% → 92.365%", "278 recovered / 37 harmful"),
        ("HiFi-CS → G1", "47.518% → 56.638%", "729 recovered / 29 harmful"),
        ("HiFi-CS → C1", "43.818% → 56.248%", "1006 recovered / 52 harmful"),
    ]
    for index, (route, score, transition) in enumerate(metrics):
        y = 2.65 + index * 1.05
        shape = slide.shapes.add_shape(1, Inches(0.75), Inches(y), Inches(11.8), Inches(0.82))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = RGBColor(218, 223, 229)
        _textbox(slide, route, 1.0, y + 0.12, 3.15, 0.45, size=21, bold=True)
        _textbox(slide, score, 4.25, y + 0.12, 3.4, 0.45, size=21, bold=True, color=BLUE)
        _textbox(slide, transition, 7.75, y + 0.14, 4.45, 0.42, size=18, color=MUTED)
    footer = "Short deck: six route cases + two identical-image cross-route comparisons" if short else "Full audit deck: one presentation-quality case for every required route/outcome pair"
    _textbox(slide, footer, 0.72, 6.45, 12.0, 0.4, size=17, bold=True, color=MAGENTA)
    _textbox(slide, "Offline 4-DoF criterion only; not physical grasp success.", 0.72, 6.92, 12.0, 0.32, size=14, color=MUTED)


def _methods_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, "How to read every case board", 0.55, 0.35, 12.2, 0.55, size=28, bold=True)
    items = [
        ("A–C", "RGB + language, evaluator GT mask, route-correct predicted-mask error overlay"),
        ("D", "Every frozen candidate in the route's All pool; no hidden below-Top-5 rows"),
        ("E", "Frozen Top-5 with cyan original Top-1 and stable native ranks"),
        ("F", "Magenta final Top-1 against one evaluator-converted matched GT grasp"),
        ("Diagnosis", "Grounding evidence → candidate generation → native rank → reranker → gate"),
        ("Criterion", "PASS iff raster IoU > 0.25 and periodic 180° angle error ≤ 30° on one GT"),
    ]
    for index, (label, text) in enumerate(items):
        y = 1.15 + index * 0.86
        _textbox(slide, label, 0.75, y, 1.7, 0.5, size=20, bold=True, color=BLUE)
        _textbox(slide, text, 2.35, y, 10.1, 0.65, size=19)
    _textbox(slide, "CROG mask quality is association-only. LightGBM contribution differences are post-hoc score decompositions, not causal proof.", 0.75, 6.45, 11.9, 0.55, size=17, bold=True, color=MAGENTA)


def _selection_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, "Selection is visual-quality optimisation within frozen outcome bins", 0.55, 0.38, 12.3, 0.65, size=27, bold=True)
    _textbox(slide, "Mandatory eligibility", 0.75, 1.35, 3.2, 0.4, size=21, bold=True, color=BLUE)
    _textbox(slide, "Verified RGB/masks; route-correct mask source; valid geometry; same GT for original/final metrics; readable target/crop/prompt; non-corrupt image.", 0.75, 1.78, 5.75, 2.0, size=19)
    _textbox(slide, "Clarity score", 6.95, 1.35, 2.7, 0.4, size=21, bold=True, color=BLUE)
    _textbox(slide, "Target visibility, crop resolution, candidate separation, matched-GT visibility, mask evidence, threshold margins, overlay legibility, prompt legibility, and mechanism purity.", 6.95, 1.78, 5.4, 2.0, size=19)
    _textbox(slide, "Anti-cherry-picking controls", 0.75, 4.35, 4.0, 0.4, size=21, bold=True, color=BLUE)
    _textbox(slide, "All eligible rows are retained; each selected case includes the next five alternatives; SHA-256 is tie-break only; the full deck covers harmful, gate and irreparable outcomes as well as recoveries.", 0.75, 4.8, 11.55, 1.2, size=19)
    _textbox(slide, "Post-formal qualitative analysis only: no training, threshold tuning, candidate change, or formal Test rerun.", 0.75, 6.5, 11.8, 0.4, size=17, bold=True, color=MAGENTA)


def _board_slide(prs: Presentation, path: Path) -> None:
    with Image.open(path) as image:
        if image.size[0] < 1920 or image.size[1] < 1080:
            raise RuntimeError(f"case-board asset resolution too small: {path}")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(path), 0, 0, width=W, height=H)


def _conclusion(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, "What the cases establish", 0.62, 0.5, 12.0, 0.6, size=30, bold=True)
    bullets = [
        "Reranking helps when a correct candidate already exists and frozen T2 evidence separates it from native Top-1.",
        "It cannot recover a sample with no correct candidate in the frozen pool; G1/C1 remain candidate-availability limited.",
        "A learned reranker can promote a formally wrong candidate; the conservative gate prevents some, but not all, harmful switches.",
        "Grounding, candidate generation, native ranking, learned reranking, and gate failures are visibly distinct operational bottlenecks.",
        "Every claim is bounded to the frozen offline 2D criterion and the available post-formal evidence.",
    ]
    for index, text in enumerate(bullets):
        _textbox(slide, f"{index + 1}", 0.8, 1.45 + index * 0.98, 0.55, 0.45, size=22, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
        _textbox(slide, text, 1.55, 1.38 + index * 0.98, 10.9, 0.72, size=20)
    _textbox(slide, "The revised gallery preserves all selected failures and their alternatives; presentation clarity does not alter the experiment.", 0.72, 6.68, 11.9, 0.42, size=17, bold=True, color=MAGENTA)


def build(run: Path) -> tuple[Path, Path]:
    run = run.resolve()
    gallery = pd.read_csv(run / "05_gallery/gallery_manifest.csv")
    cross = pd.read_csv(run / "05_gallery/cross_route_manifest.csv")
    short = pd.read_csv(run / "02_selection/SELECTED_SHORT_DECK_CASES.csv")
    full = _new()
    _title_slide(full, short=False)
    _methods_slide(full)
    _selection_slide(full)
    for route in ROUTES:
        for category in AUDIT_ORDER:
            match = gallery[(gallery["route"] == route) & (gallery["presentation_outcome"] == category)]
            if len(match) != 1:
                raise RuntimeError(f"full-deck board missing: {route}/{category}")
            _board_slide(full, Path(match.iloc[0].board_path))
    for row in cross.sort_values("cross_route_order").itertuples(index=False):
        _board_slide(full, Path(row.board_path))
    _conclusion(full)
    full_path = run / "06_slides/RERANKING_SUCCESS_FAILURE_ANALYSIS_REVISED.pptx"
    full.save(full_path)
    compact = _new()
    _title_slide(compact, short=True)
    for row in short.sort_values("short_deck_order").itertuples(index=False):
        match = gallery[(gallery["route"] == row.route) & (gallery["sample_id"] == row.sample_id)]
        if len(match) != 1:
            raise RuntimeError(f"short-deck board missing: {row.route}/{row.sample_id}")
        _board_slide(compact, Path(match.iloc[0].board_path))
    for row in cross.sort_values("cross_route_order").itertuples(index=False):
        _board_slide(compact, Path(row.board_path))
    _conclusion(compact)
    short_path = run / "06_slides/RERANKING_CASES_SHORT_DECK.pptx"
    compact.save(short_path)
    return full_path, short_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-dir", type=Path, required=True)
    args = parser.parse_args()
    for path in build(args.run_dir):
        print(path)


if __name__ == "__main__":
    main()
